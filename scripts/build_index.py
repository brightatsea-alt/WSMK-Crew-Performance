#!/usr/bin/env python3
"""
Deep-index the two WSMK workbooks into data/index.json for the Crew Performance platform.

Usage:
  python3 scripts/build_index.py "<path to WSMK PSC Status & Schedule...xlsx>" "<path to WSMK KPI.xlsx>"

Sources
  PSC  : "Database" sheet  (one row per deficiency; grouped here into inspections by vessel+date)
  UA   : "Tech Data" sheet (Unscheduled stoppage (Operational Delay) rows, Off-Hire hours)
  LTIF : "Crew Data1" sheet (LTIF / LTSF events with rank + description)
  Fleet: vessel lists (for per-vessel averages)
"""
import sys, json, re, datetime, collections
import openpyxl
import warnings
warnings.filterwarnings("ignore")

PSC_XLSX = sys.argv[1]
KPI_XLSX = sys.argv[2]
OUT = sys.argv[3] if len(sys.argv) > 3 else "data/index.json"

# ---------- helpers ----------
def norm(name):
    """Normalise vessel name: upper-case, collapse spaces, strip M/V prefixes."""
    if name is None:
        return None
    s = str(name).strip()
    s = re.sub(r"^(M/V|MV|M\.V\.)\s+", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).upper()
    return s or None

_MONTHS = {m: i for i, m in enumerate(["jan","feb","mar","apr","may","jun","jul","aug","sep","oct","nov","dec"], 1)}
def d(v):
    """Date cell -> ISO string. Accepts datetime/date, Excel serial numbers and common text formats."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:   # Excel serial
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date().isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})[.\-/]\s*(\d{1,2})[.\-/]\s*(\d{1,2})", s)              # 2026-08-09 / 2026.8.9
    if m:
        y, mo, dd = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[./\-\s]([A-Za-z]{3})[A-Za-z]*[./\-\s](\d{2,4})", s)  # 09/Aug/2026, 9 Aug 26
        if m and m.group(2).lower() in _MONTHS:
            dd, mo, y = int(m.group(1)), _MONTHS[m.group(2).lower()], int(m.group(3))
        else:
            m = re.match(r"^(\d{1,2})[./\-](\d{1,2})[./\-](\d{2,4})", s)                # 09/08/2026 (D/M/Y as used in WSMK files)
            if not m:
                return None
            dd, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    try:
        return datetime.date(y, mo, dd).isoformat()
    except ValueError:
        return None

def hours(v):
    """Convert Off-Hire cell to hours (float)."""
    if v is None:
        return 0.0
    if isinstance(v, datetime.timedelta):
        return round(v.total_seconds() / 3600, 2)
    if isinstance(v, datetime.time):
        return round(v.hour + v.minute / 60 + v.second / 3600, 2)
    if isinstance(v, datetime.datetime):
        # Excel time > 24h sometimes surfaces as datetime 1899-12-31 + days
        base = datetime.datetime(1899, 12, 30)
        return round((v - base).total_seconds() / 3600, 2)
    if isinstance(v, (int, float)):
        return round(float(v) * 24, 2)  # Excel serial fraction of a day
    s = str(v).strip().lower()
    m = re.match(r"(?:(\d+)\s*d(?:ay)?s?[,\s]*)?(\d+):(\d+)(?::(\d+))?", s)          # 26:18 / 1 day, 2:30 / 1d 2:30
    if m:
        days = int(m.group(1) or 0)
        return round(days * 24 + int(m.group(2)) + int(m.group(3)) / 60 + int(m.group(4) or 0) / 3600, 2)
    m = re.match(r"^(\d+(?:\.\d+)?)\s*h(?:r|rs|ours?)?(?:\s*(\d+)\s*m(?:in)?)?$", s)     # 5h / 5.5 hrs / 5h 30m
    if m:
        return round(float(m.group(1)) + int(m.group(2) or 0) / 60, 2)
    m = re.match(r"^(\d+(?:\.\d+)?)\s*d(?:ay)?s?$", s)                                 # 2 days
    if m:
        return round(float(m.group(1)) * 24, 2)
    try:
        return round(float(s), 2)                                                         # plain number typed as hours
    except Exception:
        return None

def clean(s, n=400):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()[:n]

# ---------- PSC ----------
wb = openpyxl.load_workbook(PSC_XLSX, read_only=True, data_only=True)
ws = wb["Database"]
insp = collections.OrderedDict()
blank = 0
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    if r[1] is None and r[4] is None:
        blank += 1
        if blank > 200:
            break
        continue
    blank = 0
    vessel = norm(r[1])
    date = d(r[4])
    if not vessel or not date:
        continue
    key = (vessel, date)
    if key not in insp:
        insp[key] = {
            "vessel": vessel, "date": date, "owner": clean(r[2], 40),
            "port": clean(r[7], 60), "country": clean(r[8], 60), "mou": clean(r[9], 40),
            "flag": clean(r[11], 10), "deficiencies": [], "detention": False,
        }
    rec = insp[key]
    code = clean(r[10], 20)
    desc = clean(r[16])
    loc = clean(r[17], 60)
    item = clean(r[19] or r[18], 80)
    if code.lower().startswith("no def") or code.lower().startswith("total") or (not code and not desc):
        continue
    if "30" in re.findall(r"\d+", code) or "detention" in desc.lower():
        rec["detention"] = True
    rec["deficiencies"].append({"code": code, "desc": desc, "location": loc, "item": item})

psc = list(insp.values())
for p in psc:
    p["def_count"] = len(p["deficiencies"])
    p["code17"] = sum(1 for x in p["deficiencies"] if "17" in re.findall(r"\d+", x["code"]))
    p["code30"] = sum(1 for x in p["deficiencies"] if "30" in re.findall(r"\d+", x["code"]))

# current fleet from schedule sheet
fleet_now = set()
try:
    ws = wb["WSMK Vessel Schedule"]
    for i, r in enumerate(ws.iter_rows(values_only=True, max_row=3000)):
        if i < 2:
            continue
        if r[0] and isinstance(r[0], str) and r[0].strip() and r[0].strip().upper() != "VESSEL NAME":
            fleet_now.add(norm(r[0]))
except Exception:
    pass

# ---------- KPI workbook: header-based column mapping (robust to inserted/moved columns) ----------
QA = {"skipped": [], "warnings": [], "counts": {}}
def find_header(rows, must=("vessel", "date", "event")):
    """Return (header_row_index, {key: col_index}) for the first row that contains all `must` labels."""
    for i, r in enumerate(rows[:15]):
        cells = [clean(c, 60).lower() for c in r]
        if all(any(k in c for c in cells) for k in must):
            col = {}
            for j, c in enumerate(cells):
                if not c: continue
                if "vessel" in c and "vessel" not in col: col["vessel"] = j
                elif c.startswith("owner") and "owner" not in col: col["owner"] = j
                elif c == "date" or c.startswith("date") and "date" not in col: col["date"] = j
                elif "event" in c and "event" not in col: col["event"] = j
                elif c.startswith("kpi") and "kpi" not in col: col["kpi"] = j
                elif ("off-hire" in c or "off hire" in c or "repair time" in c) and "hours" not in col: col["hours"] = j
                elif c.startswith("remark") and "remark" not in col: col["remark"] = j
                elif "docmap" in c and "docmap" not in col: col["docmap"] = j
                elif c.startswith("rank") and "rank" not in col: col["rank"] = j
                elif ("description" in c or c.startswith("detail")) and "desc" not in col: col["desc"] = j
            return i, col
    raise SystemExit("KPI sheet: header row with Vessel/Date/Event not found")

def kpi_flag(v):
    t = clean(v, 20)
    return "Yes" if t.lower() in ("yes", "y", "o", "true", "1") else ("No" if t.lower() in ("no", "n", "x", "false", "0") else (t or "?"))

def norm_event(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()

wb = openpyxl.load_workbook(KPI_XLSX, read_only=True, data_only=True)

# ---------- KPI: Unplanned Unavailability = Tech Data "Unscheduled stoppage (Operational Delay)" ----------
sheet = next((n for n in wb.sheetnames if n.strip().lower() == "tech data"), None) or next(n for n in wb.sheetnames if "tech" in n.lower())
rows = list(wb[sheet].iter_rows(values_only=True))
h, C = find_header(rows)
if "hours" not in C: QA["warnings"].append(f"Tech Data: Off-Hire/Repair time column not found — hours will be 0")
ua, seen = [], set()
for i, r in enumerate(rows[h + 1:], start=h + 2):
    g = lambda k: r[C[k]] if k in C and C[k] < len(r) else None
    if not g("vessel") and not g("event"):
        continue
    ev = norm_event(g("event"))
    if not re.search(r"unscheduled\s+stoppage|operational\s+delay", ev):
        continue
    vessel, date = norm(g("vessel")), d(g("date"))
    if not vessel or not date:
        QA["skipped"].append({"sheet": "Tech Data", "row": i, "reason": "vessel/date missing or unreadable", "vessel": clean(g("vessel"), 40), "date": clean(g("date"), 40)})
        continue
    hrs = hours(g("hours")) if "hours" in C else 0.0
    if hrs is None:
        QA["warnings"].append(f"Tech Data row {i} {vessel} {date}: unreadable Off-Hire value '{clean(g('hours'),40)}' → 0 h")
        hrs = 0.0
    if date > datetime.date.today().isoformat():
        QA["warnings"].append(f"Tech Data row {i} {vessel}: future date {date}")
    key = (vessel, date, round(hrs, 2), clean(g("remark"), 60).lower())
    if key in seen:
        QA["warnings"].append(f"Tech Data row {i} {vessel} {date}: duplicate row skipped")
        continue
    seen.add(key)
    ua.append({"vessel": vessel, "date": date, "hours": hrs, "remark": clean(g("remark")), "kpi": kpi_flag(g("kpi")), "docmap": clean(g("docmap"), 40)})
QA["counts"]["ua"] = len(ua)

# ---------- KPI: LTIF / LTSF = Crew Data1 ----------
sheet = next((n for n in wb.sheetnames if n.strip().lower() == "crew data1"), None) or next(n for n in wb.sheetnames if "crew data" in n.lower())
rows = list(wb[sheet].iter_rows(values_only=True))
h, C = find_header(rows)
lti = []
for i, r in enumerate(rows[h + 1:], start=h + 2):
    g = lambda k: r[C[k]] if k in C and C[k] < len(r) else None
    if not g("vessel") or not g("event"):
        continue
    ev = str(g("event"))
    vessel, date = norm(g("vessel")), d(g("date"))
    if not vessel or not date:
        QA["skipped"].append({"sheet": "Crew Data1", "row": i, "reason": "vessel/date missing or unreadable", "vessel": clean(g("vessel"), 40), "date": clean(g("date"), 40)})
        continue
    typ = "LTIF" if "LTIF" in ev.upper() else ("LTSF" if "LTSF" in ev.upper() else clean(ev, 20))
    lti.append({"vessel": vessel, "date": date, "type": typ, "rank": clean(g("rank"), 20), "desc": clean(g("desc") if "desc" in C else g("remark"), 200), "kpi": kpi_flag(g("kpi"))})
QA["counts"]["lti"] = len(lti)

# LTIF/LTSF source is Crew Data1 only (per user requirement)

# ---------- Fleet size per year ----------
vessels_by_year = collections.defaultdict(set)
for p in psc:
    vessels_by_year[int(p["date"][:4])].add(p["vessel"])
for u in ua:
    vessels_by_year[int(u["date"][:4])].add(u["vessel"])
for x in lti:
    vessels_by_year[int(x["date"][:4])].add(x["vessel"])
# HSEQ Data covers every recorded event -> good proxy of the active fleet
try:
    sheet = next(n for n in wb.sheetnames if "hseq data" in n.lower())
    rows = list(wb[sheet].iter_rows(values_only=True))
    h, C = find_header(rows, must=("vessel", "date"))
    for r in rows[h + 1:]:
        v = r[C["vessel"]] if C["vessel"] < len(r) else None
        dt = d(r[C["date"]]) if C["date"] < len(r) else None
        if v and dt:
            vessels_by_year[int(dt[:4])].add(norm(v))
except Exception as e:
    QA["warnings"].append(f"HSEQ Data fleet proxy skipped: {e}")
this_year = datetime.date.today().year
fleet = {str(y): len(v) for y, v in sorted(vessels_by_year.items())}
if fleet_now:
    fleet[str(this_year)] = max(fleet.get(str(this_year), 0), len(fleet_now))

all_vessels = sorted(set(p["vessel"] for p in psc) | set(u["vessel"] for u in ua) | set(x["vessel"] for x in lti) | fleet_now)

out = {
    "generated": datetime.datetime.now().isoformat(timespec="seconds"),
    "sources": {"psc": PSC_XLSX.split("/")[-1], "kpi": KPI_XLSX.split("/")[-1]},
    "fleet_by_year": fleet,
    "fleet_now": sorted(fleet_now),
    "vessels": all_vessels,
    "psc": sorted(psc, key=lambda x: x["date"]),
    "ua": sorted(ua, key=lambda x: x["date"]),
    "lti": sorted(lti, key=lambda x: x["date"]),
    "qa": {"ua_kpi_yes": sum(1 for u in ua if u["kpi"] == "Yes"), "ua_hours": round(sum(u["hours"] for u in ua), 1),
           "ua_vessels_not_in_fleet": sorted({u["vessel"] for u in ua} - set(all_vessels)),
           "skipped": QA["skipped"][:50], "warnings": QA["warnings"][:50]},
}
import os
os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
print(f"PSC inspections: {len(psc)}  (deficiency rows: {sum(p['def_count'] for p in psc)})")
print(f"Unscheduled stoppage events: {len(ua)}  total hours: {round(sum(u['hours'] for u in ua),1)}")
print(f"LTI/LTS events: {len(lti)}")
print(f"Fleet by year: {fleet}")
print(f"Vessels: {len(all_vessels)}  -> {OUT}")
print(f"QA: UA KPI=Yes {out['qa']['ua_kpi_yes']}/{len(ua)} · skipped rows {len(QA['skipped'])} · warnings {len(QA['warnings'])}")
for w in QA["warnings"][:20]: print("  ! " + w)
for x in QA["skipped"][:20]: print("  - skipped", x)
if out["qa"]["ua_vessels_not_in_fleet"]: print("  ! UA vessels not in vessel list:", out["qa"]["ua_vessels_not_in_fleet"])
